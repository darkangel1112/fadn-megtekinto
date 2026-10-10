from __future__ import annotations

from io import BytesIO
from pathlib import Path
import sys
import unittest

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "src"))

from app import build_targeted_export_payload
from workbook_mapper import (
    LOAN_EXPORT_ITEMS,
    LIVESTOCK_CONSUMPTION_COLUMN,
    LIVESTOCK_PRODUCT_COLUMN,
    MASTER_DATA_ROW_CODES,
    SUPPORT_EXPORT_ITEMS,
    SheetTemplate,
    build_family_member_view,
    build_general_cost_view,
    build_master_data_view,
    build_livestock_consumption_view,
    build_livestock_product_view,
    build_loan_view,
    build_support_view,
)


def _template(sheet_name: str, rows: list[dict[str, str]], codes: list[str], labels: list[str]) -> SheetTemplate:
    return SheetTemplate(
        sheet_name=sheet_name,
        title=sheet_name,
        column_codes=codes,
        column_labels=labels,
        rows=rows,
    )


def _support_templates() -> dict[str, SheetTemplate]:
    codes_by_sheet = {
        sheet_name: sorted(
            {
                row_code
                for _, row_code, source_sheet, _, _ in SUPPORT_EXPORT_ITEMS
                if row_code and source_sheet == sheet_name
            }
        )
        for sheet_name in ("t7_c", "t7_b1")
    }
    return {
        "t7_c": _template(
            "t7_c",
            [
                {"row_code": code, "row_title": f"Forrás megnevezés {code}", "dyn_row_serial": "0"}
                for code in codes_by_sheet["t7_c"]
            ],
            ["3"],
            ["támogatás összege"],
        ),
        "t7_b1": _template(
            "t7_b1",
            [
                {"row_code": code, "row_title": f"Forrás megnevezés {code}", "dyn_row_serial": "0"}
                for code in codes_by_sheet["t7_b1"]
            ],
            ["3", "4"],
            ["Egység", "eFt"],
        ),
    }


def _row(farm: str, code: str, osz: str, value: str, sheet: str) -> dict[str, str]:
    return {
        "akod": farm,
        "sor": code,
        "osz": osz,
        "ertek": value,
        "dimenzio1": "",
        "dyn_row_serial": "0",
        "sheet_name": sheet,
    }


class TargetedExportTests(unittest.TestCase):
    def setUp(self) -> None:
        self.farm = "TEST-01"
        self.templates = _support_templates()
        self.bulk_data = pd.DataFrame(
            [
                _row(self.farm, "m7094", "3", "1,25", "t7_c"),
                _row(self.farm, "m7080", "3", "1.25", "t7_c"),
                _row(self.farm, "m70801", "3", "0", "t7_c"),
                _row(self.farm, "m70802", "3", "", "t7_c"),
                _row(self.farm, "m70803", "3", "1 234,56", "t7_c"),
                _row(self.farm, "m7216", "3", "2,75", "t7_c"),
                _row(self.farm, "m7407", "3", "99", "t7_b1"),
                _row(self.farm, "m7407", "4", "2,50", "t7_b1"),
                _row(self.farm, "m0011", "3", "20 - Őstermelők családi gazdasága (ÖCsG)", "t0"),
                _row(self.farm, "m0001", "3", "2 - SZJA", "t0"),
                _row(self.farm, "m0005", "3", "4 - biogazdálkodást is folytat", "t0"),
                _row(self.farm, "m0075", "3", "2016", "t0"),
                _row(self.farm, "m0002", "3", "3 - ÁFA általános szabályai szerint adózik", "t0"),
                _row(self.farm, "m0018", "3", "2 - egyszeres könyvelés", "t0"),
                _row(self.farm, "m3610", "3", "120", "t3_b"),
                _row(self.farm, "m4082", "3", "350", "t4_a"),
                _row(self.farm, "m4355", "3", "100", "t4_a"),
                _row(self.farm, "m2201", "3", "1", "t2"),
                _row(self.farm, "m2201", "4", "120", "t2"),
                _row(self.farm, "m2201", "5", "1000", "t2"),
                _row(self.farm, "m6505", LIVESTOCK_CONSUMPTION_COLUMN, "2.5", "t6_b"),
                _row(self.farm, "m6737", LIVESTOCK_PRODUCT_COLUMN, "4.2", "t6_b"),
                _row(self.farm, "m6706", LIVESTOCK_PRODUCT_COLUMN, "1.5", "t6_b"),
                _row(self.farm, "m5411", "5", "3.5", "t5_c"),
                _row(self.farm, "m100", "3", "12.5", "t1_a"),
                _row(self.farm, "m100", "4", "0", "t1_a"),
                _row(self.farm, "m101", "3", "0", "t1_a"),
            ]
        )

    def _all_templates(self) -> dict[str, SheetTemplate]:
        templates = dict(self.templates)
        templates.update(
            {
                "t0": _template(
                    "t0",
                    [
                        {
                            "row_code": code,
                            "row_title": f"Törzsadat {code}",
                            "dyn_row_serial": "0",
                        }
                        for code in MASTER_DATA_ROW_CODES
                    ],
                    ["3"],
                    ["törzsadat"],
                ),
                "t5_c": _template(
                    "t5_c",
                    [{"row_code": "m5411", "row_title": "Őszi búza", "dyn_row_serial": "0"}],
                    ["5"],
                    ["Záróérték"],
                ),
                "t3_b": _template(
                    "t3_b",
                    [
                        {"row_code": code, "row_title": title, "dyn_row_serial": "0"}
                        for title, code, _, _ in LOAN_EXPORT_ITEMS
                    ],
                    ["3"],
                    ["Ezer Ft"],
                ),
                "t4_a": _template(
                    "t4_a",
                    [
                        {"row_code": "m4082", "row_title": "Bérmunka", "dyn_row_serial": "0"},
                        {"row_code": "m4355", "row_title": "Áramdíj", "dyn_row_serial": "0"},
                    ],
                    ["3"],
                    ["Ezer Ft"],
                ),
                "t2": _template(
                    "t2",
                    [
                        {"row_code": "m2201", "row_title": "Családtag", "dyn_row_serial": "0"},
                        {"row_code": "m2251", "row_title": "Időszaki dolgozók (családtagok)", "dyn_row_serial": "0"},
                    ],
                    ["3", "4", "5"],
                    ["létszám", "munkavégzés", "személyi jövedelem + TB"],
                ),
                "t6_b": _template(
                    "t6_b",
                    [
                        {"row_code": "m6505", "row_title": "Alomszalma", "dyn_row_serial": "0"},
                        {"row_code": "m6706", "row_title": "Étkezési tyúktojás, 1000 db", "dyn_row_serial": "0"},
                        {"row_code": "m6737", "row_title": "Trágya", "dyn_row_serial": "0"},
                    ],
                    ["8", "12"],
                    ["állattenyésztési felhasználás", "zárókészlet"],
                ),
                "t6_a": _template("t6_a", [], ["12"], ["Záróállomány"]),
                "t1_a": _template(
                    "t1_a",
                    [
                        {"row_code": "m100", "row_title": "Földterületi próbasor", "dyn_row_serial": "0"},
                        {"row_code": "m101", "row_title": "Nulla értékű próbasor", "dyn_row_serial": "0"},
                    ],
                    ["3", "4", "5", "6", "7", "8", "9", ""],
                    ["Adat 3", "Adat 4", "Adat 5", "Adat 6", "Adat 7", "Adat 8", "Adat 9", ""],
                ),
            }
        )
        return templates

    def test_master_data_view_contains_t0_rows_and_derived_akg_flag(self) -> None:
        templates = self._all_templates()

        view = build_master_data_view(self.bulk_data, templates, self.farm)

        self.assertEqual(list(view.columns), ["Üzemkód", "RowCode", "Törzsadat megnevezése", "Érték"])
        self.assertEqual(len(view), len(MASTER_DATA_ROW_CODES) + 1)
        self.assertEqual(view.iloc[0]["RowCode"], "m0011")
        self.assertEqual(view.iloc[0]["Érték"], "Őstermelők családi gazdasága (ÖCsG)")
        self.assertEqual(view.loc[view["RowCode"].eq("m0001"), "Érték"].iloc[0], "SZJA")
        self.assertEqual(view.loc[view["RowCode"].eq("m0075"), "Érték"].iloc[0], "2016")
        akg = view.loc[view["RowCode"].eq("AKG")].iloc[0]
        self.assertEqual(akg["Törzsadat megnevezése"], "AKG gazdálkodás")
        self.assertEqual(akg["Érték"], "Igen")

    def test_master_only_export_has_dedicated_filename_and_formatting(self) -> None:
        templates = self._all_templates()
        payload, filename, mime = build_targeted_export_payload(
            selected_farms=[self.farm],
            template_bundle={"sheets": templates},
            bulk_data=self.bulk_data,
            subjects=["master"],
        )

        self.assertEqual(filename, "TEST-01_torzsadatok.xlsx")
        self.assertEqual(
            mime,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook = load_workbook(BytesIO(payload), data_only=True)
        self.assertEqual(workbook.sheetnames, ["Törzsadatok"])
        worksheet = workbook["Törzsadatok"]
        self.assertEqual(worksheet.max_row, len(MASTER_DATA_ROW_CODES) + 6)
        akg_row = len(MASTER_DATA_ROW_CODES) + 2
        self.assertEqual(worksheet.cell(akg_row, 2).value, "AKG")
        self.assertEqual(worksheet.cell(akg_row, 4).value, "Igen")
        self.assertEqual(worksheet.cell(akg_row, 4).fill.fgColor.rgb, "00D9EAD3")
        self.assertEqual(worksheet.freeze_panes, "A2")
        self.assertTrue(worksheet.auto_filter.ref)
        self.assertEqual(worksheet.cell(len(MASTER_DATA_ROW_CODES) + 4, 1).value, "Hiteladatok – elérhető összesített értékek")

    def test_loan_cost_family_and_livestock_views(self) -> None:
        templates = self._all_templates()

        loans = build_loan_view(self.bulk_data, templates, self.farm)
        self.assertEqual(list(loans["RowCode"]), ["m3610"])
        self.assertEqual(loans.iloc[0]["Érték (eFt)"], 120)

        costs = build_general_cost_view(self.bulk_data, templates, self.farm)
        self.assertEqual(set(costs["RowCode"]), {"m4082", "m4355"})
        self.assertEqual(costs.loc[costs["RowCode"].eq("m4082"), "Érték (eFt)"].iloc[0], 350)

        family = build_family_member_view(self.bulk_data, templates["t2"], self.farm)
        self.assertEqual(len(family), 1)
        self.assertEqual(family.iloc[0]["RowCode"], "m2201")
        self.assertEqual(family.iloc[0]["létszám"], "1")

        livestock = build_livestock_consumption_view(self.bulk_data, templates["t6_b"], self.farm)
        self.assertEqual(len(livestock), 1)
        self.assertEqual(livestock.iloc[0]["RowCode"], "m6505")
        self.assertEqual(livestock.iloc[0]["Mennyiség"], 2.5)
        self.assertEqual(livestock.iloc[0]["Mértékegység"], "t")

        products = build_livestock_product_view(self.bulk_data, templates["t6_b"], self.farm)
        self.assertEqual(list(products["RowCode"]), ["m6706", "m6737"])
        self.assertEqual(products.iloc[0]["Mennyiség"], 1.5)
        self.assertEqual(products.iloc[0]["Mértékegység"], "1000 db")
        self.assertEqual(products.iloc[1]["Mértékegység"], "t")

    def test_new_sections_are_written_to_targeted_workbook(self) -> None:
        templates = self._all_templates()
        payload, _, _ = build_targeted_export_payload(
            selected_farms=[self.farm],
            template_bundle={"sheets": templates},
            bulk_data=self.bulk_data,
            subjects=["master", "stocks", "costs", "family"],
        )

        workbook = load_workbook(BytesIO(payload), data_only=True)
        self.assertEqual(
            workbook.sheetnames,
            ["Törzsadatok", "Készletek", "Általános költségek", "Családtagok"],
        )
        self.assertEqual(len(workbook["Törzsadatok"].tables), 1)
        stocks = workbook["Készletek"]
        self.assertEqual(len(stocks.tables), 2)
        stock_values = [cell.value for row in stocks.iter_rows() for cell in row]
        self.assertIn("Állattenyésztésben felhasznált anyagok", stock_values)
        self.assertIn("m6505", stock_values)
        self.assertIn("Állati termékek", stock_values)
        self.assertIn("m6737", stock_values)
        costs = workbook["Általános költségek"]
        self.assertIn("m4082", [cell.value for row in costs.iter_rows() for cell in row])
        family = workbook["Családtagok"]
        self.assertIn("m2201", [cell.value for row in family.iter_rows() for cell in row])

    def test_empty_livestock_consumption_has_no_empty_table_object(self) -> None:
        templates = self._all_templates()
        payload, _, _ = build_targeted_export_payload(
            selected_farms=["TEST-EMPTY"],
            template_bundle={"sheets": templates},
            bulk_data=self.bulk_data,
            subjects=["stocks"],
        )

        workbook = load_workbook(BytesIO(payload), data_only=True)
        stocks = workbook["Készletek"]
        stock_values = [cell.value for row in stocks.iter_rows() for cell in row]
        self.assertIn("Állattenyésztésben felhasznált anyagok", stock_values)
        self.assertIn("Állati termékek", stock_values)
        self.assertEqual(len(stocks.tables), 0)

    def test_support_list_keeps_all_rows_and_distinguishes_source_states(self) -> None:
        view = build_support_view(self.bulk_data, self.templates, self.farm)

        self.assertEqual(len(view), len(SUPPORT_EXPORT_ITEMS))
        self.assertEqual(view.iloc[0]["FADN-sorkód"], "m7094")
        self.assertEqual(view.iloc[0]["Összeg (Ft)"], 1250)
        self.assertEqual(view.iloc[0]["Sor jellege"], "Összesítő")

        zero = view.loc[view["FADN-sorkód"].eq("m70801")].iloc[0]
        self.assertEqual(zero["Összeg (Ft)"], 0)
        self.assertEqual(zero["Adatállapot"], "Forrásban 0")

        blank = view.loc[view["FADN-sorkód"].eq("m70802")].iloc[0]
        self.assertEqual(blank["Összeg (Ft)"], 0)
        self.assertEqual(blank["Adatállapot"], "Nincs kitöltött forrásérték")

        converted = view.loc[view["FADN-sorkód"].eq("m70803")].iloc[0]
        self.assertEqual(converted["Összeg (Ft)"], 1_234_560)

        cis_yf = view.loc[view["FADN-sorkód"].eq("m7407")].iloc[0]
        self.assertEqual(cis_yf["Összeg (Ft)"], 2500)

        smallholder = view.loc[view["Célűrlap jogcíme"].eq("Kistermelői támogatási rendszer")].iloc[0]
        self.assertTrue(pd.isna(smallholder["Összeg (Ft)"]))
        self.assertEqual(smallholder["Adatállapot"], "Nincs azonosított FADN-sorkód")

    def test_export_contains_support_sheet_after_land_and_preserves_land_export(self) -> None:
        templates = self._all_templates()
        payload, filename, mime = build_targeted_export_payload(
            selected_farms=[self.farm],
            template_bundle={"sheets": templates},
            bulk_data=self.bulk_data,
            subjects=["autumn", "stocks", "animals", "land", "support"],
        )

        self.assertEqual(filename, "TEST-01_celzott_export.xlsx")
        self.assertEqual(
            mime,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook = load_workbook(BytesIO(payload), data_only=True)
        self.assertEqual(
            workbook.sheetnames,
            [
                "Ősszel vetett terület",
                "Készletek",
                "Állatok",
                "Földterületi adatok",
                "Támogatási jogcímek",
            ],
        )

        land = workbook["Földterületi adatok"]
        self.assertEqual(land.max_row, 2)
        self.assertEqual(land.cell(2, 2).value, "m100")
        self.assertEqual(land.cell(2, 4).value, 12.5)

        support = workbook["Támogatási jogcímek"]
        self.assertEqual(support.max_row, len(SUPPORT_EXPORT_ITEMS) + 1)
        self.assertEqual(support.cell(2, 3).value, "m7094")
        self.assertEqual(support.cell(2, 5).value, 1250)
        self.assertEqual(support.cell(2, 5).number_format, "#,##0")
        self.assertEqual(support.freeze_panes, "A2")
        self.assertTrue(support.auto_filter.ref)

    def test_support_only_export_has_dedicated_filename(self) -> None:
        payload, filename, _ = build_targeted_export_payload(
            selected_farms=[self.farm],
            template_bundle={"sheets": self.templates},
            bulk_data=self.bulk_data,
            subjects=["support"],
        )

        self.assertEqual(filename, "TEST-01_tamogatasi_jogcimek.xlsx")
        workbook = load_workbook(BytesIO(payload), data_only=True)
        self.assertEqual(workbook.sheetnames, ["Támogatási jogcímek"])


if __name__ == "__main__":
    unittest.main()
