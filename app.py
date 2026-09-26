from __future__ import annotations

from datetime import datetime
import hashlib
from io import BytesIO
import re
import sys
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pandas as pd
import streamlit as st
import streamlit.components.v1 as components
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

APP_DIR = Path(__file__).resolve().parent
SRC_DIR = APP_DIR / "src"
HELP_FILE = APP_DIR / "sugo.md"
WEEKDAYS_HU = {
    0: "Hétfő",
    1: "Kedd",
    2: "Szerda",
    3: "Csütörtök",
    4: "Péntek",
    5: "Szombat",
    6: "Vasárnap",
}
if str(SRC_DIR) not in sys.path:
    sys.path.insert(0, str(SRC_DIR))

from workbook_mapper import (
    attach_sheet_names,
    build_animal_view,
    build_autumn_sown_view,
    build_land_view,
    build_sheet_view,
    build_stock_view,
    list_farms,
    load_bulk_data,
    load_template,
    normalize_display_value,
    summarize_template,
)


st.set_page_config(
    page_title="FADN Megtekintő",
    page_icon="🧪",
    layout="wide",
    initial_sidebar_state="expanded",
)


def inject_styles() -> None:
    st.markdown(
        """
        <style>
        :root {
            --bg-main: #071411;
            --bg-panel: rgba(7, 32, 28, 0.88);
            --line: rgba(0, 255, 208, 0.18);
            --text-main: #dbfff8;
            --text-soft: #89d8ca;
            --accent: #00f7c2;
            --accent-2: #3dd9ff;
            --accent-3: #95ff62;
        }
        .stApp {
            background:
                radial-gradient(circle at top right, rgba(61, 217, 255, 0.14), transparent 30%),
                radial-gradient(circle at top left, rgba(0, 247, 194, 0.12), transparent 28%),
                linear-gradient(180deg, #05110f 0%, #081916 45%, #06100f 100%);
            color: var(--text-main);
        }
        [data-testid="stSidebar"] {
            background:
                linear-gradient(180deg, rgba(4, 20, 17, 0.98), rgba(5, 30, 26, 0.96));
            border-right: 1px solid var(--line);
        }
        .block-container {
            padding-top: 2rem;
            padding-bottom: 2rem;
        }
        .hero {
            padding: 1.3rem 1.4rem;
            border: 1px solid var(--line);
            border-radius: 18px;
            background: linear-gradient(135deg, rgba(0, 247, 194, 0.08), rgba(61, 217, 255, 0.06));
            box-shadow: 0 0 40px rgba(0, 247, 194, 0.06);
            margin-bottom: 1rem;
        }
        .hero h1 {
            margin: 0;
            font-size: 2rem;
            color: var(--text-main);
            letter-spacing: 0.02em;
        }
        .hero p {
            margin: 0.55rem 0 0;
            color: var(--text-soft);
        }
        .hero-warning {
            margin-top: 0.85rem;
            padding: 0.8rem 0.95rem;
            border-radius: 14px;
            border: 1px solid rgba(255, 215, 0, 0.18);
            background: rgba(255, 215, 0, 0.06);
            color: #fff0a8;
            line-height: 1.45;
        }
        .chip {
            display: inline-block;
            padding: 0.3rem 0.65rem;
            margin-right: 0.45rem;
            border-radius: 999px;
            border: 1px solid rgba(149, 255, 98, 0.28);
            color: #d6ffbf;
            background: rgba(149, 255, 98, 0.08);
            font-size: 0.84rem;
        }
        .stMetric {
            background: rgba(5, 28, 24, 0.72);
            border: 1px solid var(--line);
            padding: 0.75rem;
            border-radius: 16px;
        }
        .stDataFrame, .stTable {
            border: 1px solid var(--line);
            border-radius: 14px;
            overflow: hidden;
        }
        div[data-baseweb="select"] > div,
        div[data-baseweb="input"] > div,
        .stTextInput > div > div,
        .stNumberInput > div > div {
            background: rgba(5, 28, 24, 0.9);
            border-color: rgba(0, 247, 194, 0.22);
        }
        div[data-baseweb="menu"] [role="option"] {
            cursor: pointer !important;
        }
        .stButton > button, .stDownloadButton > button {
            background: linear-gradient(90deg, rgba(0, 247, 194, 0.18), rgba(61, 217, 255, 0.16));
            color: var(--text-main);
            border: 1px solid rgba(61, 217, 255, 0.25);
            border-radius: 12px;
        }
        .hero-side {
            display: flex;
            flex-direction: column;
            align-items: flex-end;
            gap: 0.7rem;
        }
        .hero-clock {
            width: 100%;
            max-width: 220px;
            padding: 0.9rem 1rem;
            border-radius: 16px;
            border: 1px solid rgba(61, 217, 255, 0.22);
            background: rgba(5, 28, 24, 0.7);
            text-align: center;
        }
        .hero-clock-head {
            display: flex;
            justify-content: center;
            align-items: baseline;
            gap: 0.45rem;
            flex-wrap: wrap;
        }
        .hero-clock-date {
            color: var(--text-main);
            font-size: 1.02rem;
            font-weight: 600;
        }
        .hero-clock-day {
            color: #9cf6df;
            font-size: 0.98rem;
        }
        .hero-clock-time {
            color: #d6ffbf;
            margin-top: 0.45rem;
            font-size: 1.35rem;
            font-weight: 700;
            letter-spacing: 0.04em;
        }
        .hero-clock-author {
            margin-top: 0.5rem;
            color: var(--text-soft);
            font-size: 0.82rem;
            line-height: 1.35;
        }
        </style>
        """,
        unsafe_allow_html=True,
    )


def inject_language_hints() -> None:
    components.html(
        """
        <script>
        (function () {
            const doc = window.parent.document;
            const root = doc.documentElement;
            const body = doc.body;
            root.lang = "hu";
            root.setAttribute("translate", "no");
            root.classList.add("notranslate");
            if (body) {
                body.lang = "hu";
                body.setAttribute("translate", "no");
                body.classList.add("notranslate");
            }

            const head = doc.head;

            function ensureMeta(name, content) {
                let meta = head.querySelector(`meta[name="${name}"]`);
                if (!meta) {
                    meta = doc.createElement("meta");
                    meta.setAttribute("name", name);
                    head.appendChild(meta);
                }
                meta.setAttribute("content", content);
            }

            ensureMeta("google", "notranslate");
            ensureMeta("content-language", "hu");

            const targets = doc.querySelectorAll(
                ".stApp, .main, [data-testid='stAppViewContainer'], [data-testid='stSidebar'], .block-container"
            );
            targets.forEach((node) => {
                node.setAttribute("translate", "no");
                node.classList.add("notranslate");
                node.lang = "hu";
            });
        })();
        </script>
        """,
        height=0,
        width=0,
    )


def clear_targeted_export_state() -> None:
    for key in [
        "target_export_download_data",
        "target_export_download_name",
        "target_export_download_mime",
        "target_export_error",
    ]:
        st.session_state.pop(key, None)
    st.session_state["target_export_open"] = False
    st.session_state["target_export_generating"] = False


def open_targeted_export(farms: list[str]) -> None:
    st.session_state["target_export_open"] = True
    st.session_state["target_export_generating"] = False
    st.session_state["target_export_subject_autumn"] = True
    st.session_state["target_export_subject_stocks"] = True
    st.session_state["target_export_subject_animals"] = True
    st.session_state["target_export_subject_land"] = True
    for farm_code in farms:
        st.session_state[_target_farm_key(farm_code)] = False
    clear_targeted_export_download()


def clear_targeted_export_download() -> None:
    for key in [
        "target_export_download_data",
        "target_export_download_name",
        "target_export_download_mime",
        "target_export_error",
    ]:
        st.session_state.pop(key, None)


def clear_uploaded_sources() -> None:
    st.session_state["uploader_nonce"] = st.session_state.get("uploader_nonce", 0) + 1
    for key in [
        "loaded_source_signature",
        "loaded_bulk_data",
        "loaded_template_bundle",
        "loaded_bulk_name",
        "loaded_template_name",
    ]:
        st.session_state.pop(key, None)
    clear_targeted_export_state()


def read_help_text() -> str:
    if HELP_FILE.exists():
        return HELP_FILE.read_text(encoding="utf-8").strip()
    return "A súgó szövege még nincs megadva."


@st.dialog("Súgó", width="large")
def show_help_dialog() -> None:
    with st.container(height=620, border=False):
        st.markdown(read_help_text())
    if st.button("Bezárás", use_container_width=True):
        st.rerun()


def _source_signature(bulk_bytes: bytes, template_bytes: bytes) -> str:
    digest = hashlib.sha1()
    digest.update(bulk_bytes)
    digest.update(b"::")
    digest.update(template_bytes)
    return digest.hexdigest()


def load_session_sources(
    bulk_bytes: bytes,
    template_bytes: bytes,
    bulk_name: str,
    template_name: str,
) -> tuple[dict, pd.DataFrame]:
    signature = _source_signature(bulk_bytes, template_bytes)
    if st.session_state.get("loaded_source_signature") != signature:
        template_bundle = load_template(template_bytes)
        bulk_data = load_bulk_data(bulk_bytes)
        bulk_data = attach_sheet_names(bulk_data, template_bundle["row_to_sheet"])

        st.session_state["loaded_source_signature"] = signature
        st.session_state["loaded_template_bundle"] = template_bundle
        st.session_state["loaded_bulk_data"] = bulk_data
        st.session_state["loaded_bulk_name"] = bulk_name
        st.session_state["loaded_template_name"] = template_name
        clear_targeted_export_state()

    return (
        st.session_state["loaded_template_bundle"],
        st.session_state["loaded_bulk_data"],
    )


def resolve_sources():
    if "uploader_nonce" not in st.session_state:
        st.session_state["uploader_nonce"] = 0

    st.sidebar.subheader("Adatforrások")
    st.sidebar.caption("Az app csak a feltöltött fájlokból dolgozik.")

    bulk_upload = st.sidebar.file_uploader(
        "Ömlesztett Excel",
        type=["xlsx", "xlsm", "xls"],
        key=f"bulk_upload_{st.session_state['uploader_nonce']}",
    )
    template_upload = st.sidebar.file_uploader(
        "Táblázatos sablon Excel",
        type=["xlsx", "xlsm", "xls"],
        key=f"template_upload_{st.session_state['uploader_nonce']}",
    )

    st.sidebar.button(
        "Források törlése",
        on_click=clear_uploaded_sources,
        use_container_width=True,
    )
    if bulk_upload is None or template_upload is None:
        st.sidebar.button(
            "Célzott export",
            disabled=True,
            use_container_width=True,
        )
        st.info("A kezdéshez tölts fel egy ömlesztett és egy táblázatos Excel-fájlt a bal oldali sávban.")
        st.stop()

    bulk_bytes = bulk_upload.getvalue()
    template_bytes = template_upload.getvalue()
    bulk_name = bulk_upload.name
    template_name = template_upload.name
    return bulk_bytes, template_bytes, bulk_name, template_name


def render_header() -> None:
    now = datetime.now()
    date_text = now.strftime("%Y.%m.%d.")
    day_text = WEEKDAYS_HU[now.weekday()]
    time_text = now.strftime("%H:%M")

    left_col, right_col = st.columns([6, 1.6], gap="large")

    with left_col:
        st.markdown(
            """
            <div class="hero">
                <div class="chip">Alapverzió</div>
                <div class="chip">Ömlesztett -> Táblázatos nézet</div>
                <div class="chip">Streamlit alap</div>
                <h1>FADN Megtekintő</h1>
                <p>
                    Válassz üzemet és munkalapot, az app pedig a sablonmunkafüzet
                    alapján olvasható táblázatot épít az ömlesztett adatokból.
                </p>
                <div class="hero-warning">
                    Figyelem! A feltöltött file-ok a szerveren nem kerülnek tárolásra,
                    az ablak bezárásával minden adat törlődik.
                </div>
            </div>
            """,
            unsafe_allow_html=True,
        )

    with right_col:
        st.markdown(
            f"""
            <div class="hero hero-side">
                <div class="hero-clock">
                    <div class="hero-clock-head">
                        <div class="hero-clock-date">{date_text}</div>
                        <div class="hero-clock-day">{day_text}</div>
                    </div>
                    <div class="hero-clock-time">{time_text}</div>
                    <div class="hero-clock-author">Készítette: Kovács Péter DDRI</div>
                </div>
            </div>
            """,
            unsafe_allow_html=True,
        )
        if st.button("Súgó", use_container_width=True):
            show_help_dialog()


def _sheet_select_label(sheet_key: str, sheet_title: str, max_length: int = 38) -> str:
    """Keep long worksheet names from reflowing the sidebar select menu."""

    normalized_title = re.sub(r"\s+", " ", str(sheet_title)).strip()
    full_label = f"{sheet_key} - {normalized_title}"
    if len(full_label) <= max_length:
        return full_label
    return f"{full_label[: max_length - 1].rstrip()}…"


def dataframe_to_csv_bytes(dataframe: pd.DataFrame) -> bytes:
    return dataframe.to_csv(index=False).encode("utf-8-sig")


def _safe_farm_filename(farm_code: str) -> str:
    safe_name = re.sub(r"[^0-9A-Za-z_-]+", "-", farm_code).strip("-")
    return safe_name or "uzem"


def dataframes_to_excel_bytes(
    sheet_frames: list[tuple[str, pd.DataFrame]],
) -> bytes:
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for sheet_name, dataframe in sheet_frames:
            dataframe.to_excel(writer, index=False, sheet_name=sheet_name)
            worksheet = writer.sheets[sheet_name]
            worksheet.freeze_panes = "A2"
            has_data_rows = not dataframe.empty
            if has_data_rows and sheet_name != "Állatok":
                worksheet.auto_filter.ref = worksheet.dimensions
            worksheet.sheet_view.showGridLines = True
            worksheet.sheet_properties.tabColor = "00F7C2"

            header_fill = PatternFill(fill_type="solid", fgColor="0B4A40")
            header_font = Font(bold=True, color="DBFFF8")
            thin_border = Border(
                left=Side(style="thin", color="B7C9C5"),
                right=Side(style="thin", color="B7C9C5"),
                top=Side(style="thin", color="B7C9C5"),
                bottom=Side(style="thin", color="B7C9C5"),
            )
            for row in worksheet.iter_rows():
                for cell in row:
                    cell.border = thin_border

            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            for column_cells in worksheet.columns:
                column_letter = column_cells[0].column_letter
                values = [str(cell.value or "") for cell in column_cells]
                width = min(max(max(len(value) for value in values) + 2, 12), 42)
                worksheet.column_dimensions[column_letter].width = width

            numeric_columns = {
                column
                for column in dataframe.columns
                if pd.api.types.is_numeric_dtype(dataframe[column])
                or column in {
                    "Érték",
                    "Záróérték",
                    "5C záróérték",
                    "6B zárókészlet",
                    "Saját készlet",
                    "Vásárolt készlet",
                }
            }
            for column in numeric_columns:
                value_column_index = list(dataframe.columns).index(column) + 1
                for row in worksheet.iter_rows(
                    min_row=2,
                    min_col=value_column_index,
                    max_col=value_column_index,
                ):
                    row[0].number_format = "0.00"

            if "Ellenőrzés" in dataframe.columns:
                status_column_index = list(dataframe.columns).index("Ellenőrzés") + 1
                for row in worksheet.iter_rows(
                    min_row=2,
                    min_col=status_column_index,
                    max_col=status_column_index,
                ):
                    cell = row[0]
                    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                    if str(cell.value or "").startswith("HIBA"):
                        cell.fill = PatternFill(fill_type="solid", fgColor="F4CCCC")
                        cell.font = Font(color="9C0006", bold=True)
                    else:
                        cell.fill = PatternFill(fill_type="solid", fgColor="D9EAD3")
                        cell.font = Font(color="274E13")

            if sheet_name == "Készletek":
                worksheet.column_dimensions["C"].width = 34
                worksheet.column_dimensions["F"].width = 34
                worksheet.column_dimensions["J"].width = 48
            elif sheet_name == "Állatok":
                worksheet.column_dimensions["C"].width = 42
                worksheet.column_dimensions["D"].width = 16
                if has_data_rows:
                    animal_table = Table(displayName="tbl_Allatok", ref=worksheet.dimensions)
                    animal_table.tableStyleInfo = TableStyleInfo(
                        name="TableStyleMedium2",
                        showFirstColumn=False,
                        showLastColumn=False,
                        showRowStripes=False,
                        showColumnStripes=False,
                    )
                    worksheet.add_table(animal_table)

    return output.getvalue()


def dataframe_to_excel_bytes(dataframe: pd.DataFrame, sheet_name: str) -> bytes:
    return dataframes_to_excel_bytes([(sheet_name, dataframe)])


def _targeted_export_filename(farm_code: str, subjects: list[str]) -> str:
    safe_farm = _safe_farm_filename(farm_code)
    subject_set = set(subjects)
    if subject_set == {"autumn"}:
        return f"{safe_farm}_ossszel_vetett_terulet.xlsx"
    if subject_set == {"stocks"}:
        return f"{safe_farm}_keszletek.xlsx"
    if subject_set == {"animals"}:
        return f"{safe_farm}_allatok.xlsx"
    if subject_set == {"land"}:
        return f"{safe_farm}_foldteruleti_adatok.xlsx"
    return f"{safe_farm}_celzott_export.xlsx"


def build_targeted_export_payload(
    selected_farms: list[str],
    template_bundle: dict,
    bulk_data: pd.DataFrame,
    subjects: list[str],
) -> tuple[bytes, str, str]:
    export_files: list[tuple[str, bytes]] = []

    for farm_code in selected_farms:
        sheet_frames: list[tuple[str, pd.DataFrame]] = []
        if "autumn" in subjects:
            autumn_template = template_bundle["sheets"]["t5_c"]
            sheet_frames.append(
                (
                    "Ősszel vetett terület",
                    build_autumn_sown_view(
                        dataframe=bulk_data,
                        template=autumn_template,
                        farm_code=farm_code,
                    ),
                )
            )

        if "stocks" in subjects:
            sheet_frames.append(
                (
                    "Készletek",
                    build_stock_view(
                        dataframe=bulk_data,
                        template_5c=template_bundle["sheets"]["t5_c"],
                        template_6b=template_bundle["sheets"]["t6_b"],
                        farm_code=farm_code,
                    ),
                )
            )

        if "animals" in subjects:
            sheet_frames.append(
                (
                    "Állatok",
                    build_animal_view(
                        dataframe=bulk_data,
                        template=template_bundle["sheets"]["t6_a"],
                        farm_code=farm_code,
                    ),
                )
            )

        if "land" in subjects:
            sheet_frames.append(
                (
                    "Földterületi adatok",
                    build_land_view(
                        dataframe=bulk_data,
                        template=template_bundle["sheets"]["t1_a"],
                        farm_code=farm_code,
                    ),
                )
            )

        workbook_bytes = dataframes_to_excel_bytes(sheet_frames)
        export_files.append(
            (
                _targeted_export_filename(farm_code, subjects),
                workbook_bytes,
            )
        )

    if len(export_files) == 1:
        filename, workbook_bytes = export_files[0]
        return workbook_bytes, filename, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    archive_buffer = BytesIO()
    with ZipFile(archive_buffer, mode="w", compression=ZIP_DEFLATED) as archive:
        for filename, workbook_bytes in export_files:
            archive.writestr(filename, workbook_bytes)

    return archive_buffer.getvalue(), "celzott_export.zip", "application/zip"


def _target_farm_key(farm_code: str) -> str:
    return f"target_farm_{farm_code}"


def _initialize_target_farm_state(farms: list[str]) -> None:
    farm_keys = {_target_farm_key(farm_code) for farm_code in farms}
    for key in list(st.session_state.keys()):
        if key.startswith("target_farm_") and key not in farm_keys:
            st.session_state.pop(key, None)
    for farm_code in farms:
        st.session_state.setdefault(_target_farm_key(farm_code), False)


def _select_all_target_farms(farms: list[str]) -> None:
    for farm_code in farms:
        st.session_state[_target_farm_key(farm_code)] = True
    clear_targeted_export_download()


def _clear_all_target_farms(farms: list[str]) -> None:
    for farm_code in farms:
        st.session_state[_target_farm_key(farm_code)] = False
    clear_targeted_export_download()


@st.dialog(
    "Célzott export",
    width="large",
    on_dismiss=clear_targeted_export_state,
)
def show_targeted_export_dialog(
    farms: list[str],
    template_bundle: dict,
    bulk_data: pd.DataFrame,
) -> None:
    _initialize_target_farm_state(farms)
    st.session_state.setdefault("target_export_subject_autumn", True)
    st.session_state.setdefault("target_export_subject_stocks", True)
    st.session_state.setdefault("target_export_subject_animals", True)
    st.session_state.setdefault("target_export_subject_land", True)
    st.session_state.setdefault("target_export_generating", False)

    st.caption("Válaszd ki az üzemeket és az export tárgyát.")
    farm_column, subject_column = st.columns([1.45, 1])

    with farm_column:
        st.subheader("Üzemek")
        select_all_column, clear_all_column = st.columns(2)
        with select_all_column:
            st.button(
                "Összes kijelölése",
                on_click=_select_all_target_farms,
                args=(farms,),
                use_container_width=True,
            )
        with clear_all_column:
            st.button(
                "Kijelölés törlése",
                on_click=_clear_all_target_farms,
                args=(farms,),
                use_container_width=True,
            )

        with st.container(height=360, border=True):
            for farm_code in farms:
                st.checkbox(
                    farm_code,
                    key=_target_farm_key(farm_code),
                    on_change=clear_targeted_export_download,
                )

    with subject_column:
        st.subheader("Export tárgya")
        autumn_available = "t5_c" in template_bundle["sheets"]
        stocks_available = "t5_c" in template_bundle["sheets"] and "t6_b" in template_bundle["sheets"]
        animals_available = "t6_a" in template_bundle["sheets"]
        land_available = "t1_a" in template_bundle["sheets"]
        st.checkbox(
            "Ősszel vetett terület",
            key="target_export_subject_autumn",
            disabled=not autumn_available,
            on_change=clear_targeted_export_download,
            help=(
                "Az 5C mezei leltárának nem nulla záróértékű vetési sorai kerülnek az "
                "Ősszel vetett terület munkalapra. A búza sorai külön is megmaradnak, "
                "és szükség esetén külön összesített Búza sor is készül."
                if autumn_available
                else "Az őszi vetett terület exportjához az 5C munkalapnak is be kell töltődnie."
            ),
        )
        st.checkbox(
            "Készletek",
            key="target_export_subject_stocks",
            disabled=not stocks_available,
            on_change=clear_targeted_export_download,
            help=(
                "A készletek exportja az 5C és 6B munkalapból készül."
                if stocks_available
                else "A készletek exportjához az 5C és 6B munkalapnak is be kell töltődnie."
            ),
        )
        st.checkbox(
            "Állatok",
            key="target_export_subject_animals",
            disabled=not animals_available,
            on_change=clear_targeted_export_download,
            help=(
                "A 6A munkalap nem nulla záróállományai kerülnek az Állatok "
                "munkalapra. Az összesítő sorok kimaradnak, a nem nulla, "
                "tonnában megadott súlysorok megmaradnak."
                if animals_available
                else "Az állatok exportjához a 6A munkalapnak is be kell töltődnie."
            ),
        )
        st.checkbox(
            "Földterületi adatok",
            key="target_export_subject_land",
            disabled=not land_available,
            on_change=clear_targeted_export_download,
            help=(
                "A t1_a munkalap minden olyan sorát exportálja, amelyben legalább egy "
                "nem nulla érték található. Az értékoszlopok a megjelenítő sorrendjét követik."
                if land_available
                else "A földterületi exporthoz a t1_a munkalapnak is be kell töltődnie."
            ),
        )

        selected_farms = [
            farm_code
            for farm_code in farms
            if st.session_state.get(_target_farm_key(farm_code), False)
        ]
        selected_subjects = []
        if st.session_state["target_export_subject_autumn"] and autumn_available:
            selected_subjects.append("autumn")
        if st.session_state["target_export_subject_stocks"] and stocks_available:
            selected_subjects.append("stocks")
        if st.session_state["target_export_subject_animals"] and animals_available:
            selected_subjects.append("animals")
        if st.session_state["target_export_subject_land"] and land_available:
            selected_subjects.append("land")
        st.caption(f"Kijelölt üzemek: {len(selected_farms)}")

        export_enabled = bool(selected_farms) and bool(selected_subjects)
        if st.button(
            "Export létrehozása",
            type="primary",
            use_container_width=True,
            disabled=not export_enabled or st.session_state["target_export_generating"],
        ):
            clear_targeted_export_download()
            st.session_state["target_export_generating"] = True
            st.rerun()

        if st.session_state.get("target_export_generating", False):
            try:
                with st.spinner("Az export készül, kérlek várj..."):
                    payload, filename, mime = build_targeted_export_payload(
                        selected_farms=selected_farms,
                        template_bundle=template_bundle,
                        bulk_data=bulk_data,
                        subjects=selected_subjects,
                    )
            except ValueError as error:
                st.session_state["target_export_generating"] = False
                st.session_state["target_export_error"] = str(error)
            else:
                st.session_state["target_export_download_data"] = payload
                st.session_state["target_export_download_name"] = filename
                st.session_state["target_export_download_mime"] = mime
                st.session_state["target_export_generating"] = False
                st.success("Az export elkészült, letölthető.")

        if st.session_state.get("target_export_error"):
            st.error(st.session_state["target_export_error"])

        download_data = st.session_state.get("target_export_download_data")
        if download_data:
            st.download_button(
                "Export letöltése",
                data=download_data,
                file_name=st.session_state["target_export_download_name"],
                mime=st.session_state["target_export_download_mime"],
                use_container_width=True,
            )

    st.divider()
    if st.button("Bezárás", use_container_width=True):
        clear_targeted_export_state()
        st.rerun(scope="app")


def main() -> None:
    inject_styles()
    inject_language_hints()
    render_header()

    bulk_bytes, template_bytes, bulk_name, template_name = resolve_sources()

    with st.spinner("Munkafüzetek beolvasása..."):
        template_bundle, bulk_data = load_session_sources(
            bulk_bytes=bulk_bytes,
            template_bytes=template_bytes,
            bulk_name=bulk_name,
            template_name=template_name,
        )

    farms = list_farms(bulk_data)
    sheets = template_bundle["sheets"]
    template_summary = summarize_template(template_bundle)

    target_export_available = bool(bulk_data.shape[0]) and any(
        sheet_name in sheets for sheet_name in ("t5_c", "t6_b", "t6_a", "t1_a")
    )
    target_export_requested = st.sidebar.button(
        "Célzott export",
        disabled=not target_export_available,
        use_container_width=True,
    )

    st.sidebar.subheader("Nézet")
    farm_code = st.sidebar.selectbox(
        "Üzem",
        farms,
        index=0 if farms else None,
        key="main_farm_code",
    )
    selected_sheet_name = st.sidebar.selectbox(
        "Munkalap",
        options=list(sheets.keys()),
        format_func=lambda key: _sheet_select_label(key, sheets[key].title),
        key="main_selected_sheet",
    )
    filled_only = st.sidebar.checkbox(
        "Csak kitöltött sorok",
        value=False,
        key="main_filled_only",
    )
    closing_only = st.sidebar.checkbox(
        "Kísérleti: csak záró/összesítő sorok",
        value=False,
        key="main_closing_only",
    )
    show_hidden_technical = st.sidebar.checkbox(
        "Rejtett technikai oszlopok megnyitása",
        value=False,
        key="main_show_hidden_technical",
    )

    unknown_count = int((bulk_data["sheet_name"] == "").sum())
    if unknown_count:
        st.sidebar.warning(f"{unknown_count} rekordhoz nem találtam sablonmunkalapot.")

    current_sheet = sheets[selected_sheet_name]
    view_df = build_sheet_view(
        dataframe=bulk_data,
        template=current_sheet,
        farm_code=farm_code,
        filled_only=filled_only,
        closing_only=closing_only,
    )
    full_view_df = build_sheet_view(
        dataframe=bulk_data,
        template=current_sheet,
        farm_code=farm_code,
        filled_only=False,
        closing_only=closing_only,
    )

    metric_col_1, metric_col_2, metric_col_3, metric_col_4 = st.columns(4)
    metric_col_1.metric("Üzemek száma", len(farms))
    metric_col_2.metric("Sablonlapok", len(sheets))
    metric_col_3.metric("Ömlesztett rekordok", len(bulk_data))
    metric_col_4.metric("Aktuális nézet sorai", len(view_df))

    st.caption(f"Források: `{bulk_name}` és `{template_name}`")

    tab_view, tab_mapping, tab_raw = st.tabs(["Nézet", "Leképzés", "Nyers adatok"])

    with tab_view:
        st.subheader(current_sheet.title)
        st.caption(
            f"Munkalap azonosító: `{current_sheet.sheet_name}`. "
            f"Értékoszlopok: {len(current_sheet.column_codes)}."
        )
        if filled_only:
            hidden_rows = len(full_view_df) - len(view_df)
            if hidden_rows == 0:
                st.info("Ezen a lapon most minden sorban van érdemi adat, ezért a szűrés nem változtat a listán.")
            else:
                st.caption(f"A szűrés {hidden_rows} üres vagy csak technikai `= 0` értéket tartalmazó sort rejtett el.")

        display_df = view_df.copy()
        display_df = display_df.map(normalize_display_value)

        preferred_order = [
            column
            for column in ["FarmCode", "RowCode", "RowTitle"]
            if column in display_df.columns
        ]
        value_columns = [
            column
            for column in display_df.columns
            if column not in ["FarmCode", "RowCode", "RowTitle", "DynRowSerial", "Dimension1"]
        ]
        technical_tail = []
        if show_hidden_technical:
            technical_tail = [
                column
                for column in ["DynRowSerial", "Dimension1"]
                if column in display_df.columns
            ]

        display_df = display_df[preferred_order + value_columns + technical_tail]

        st.dataframe(display_df, use_container_width=True, height=680)
        st.download_button(
            "Aktuális nézet letöltése CSV-ben",
            data=dataframe_to_csv_bytes(display_df),
            file_name=f"{farm_code}_{current_sheet.sheet_name}.csv",
            mime="text/csv",
        )

    with tab_mapping:
        st.subheader("Sablonösszefoglaló")
        st.dataframe(template_summary, use_container_width=True, height=420)

        st.subheader("Aktuális lap oszloptérképe")
        column_map_df = pd.DataFrame(
            {
                "Oszlop kód": current_sheet.column_codes,
                "Megjelenített címke": current_sheet.column_labels,
            }
        )
        st.dataframe(column_map_df, use_container_width=True, height=320)

    with tab_raw:
        st.subheader("Nyers ömlesztett rekordok az aktuális laphoz")
        raw_df = bulk_data[
            (bulk_data["akod"] == farm_code) & (bulk_data["sheet_name"] == current_sheet.sheet_name)
        ].copy()
        st.dataframe(raw_df, use_container_width=True, height=620)

    if target_export_requested:
        open_targeted_export(farms)
        show_targeted_export_dialog(
            farms=farms,
            template_bundle=template_bundle,
            bulk_data=bulk_data,
        )
    elif st.session_state.get("target_export_open", False):
        show_targeted_export_dialog(
            farms=farms,
            template_bundle=template_bundle,
            bulk_data=bulk_data,
        )


if __name__ == "__main__":
    main()
