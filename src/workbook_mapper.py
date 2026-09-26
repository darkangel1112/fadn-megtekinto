from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any
import re
import zipfile
import xml.etree.ElementTree as ET

import pandas as pd
from openpyxl import load_workbook


TECHNICAL_COLUMNS = ["FarmCode", "RowCode", "RowTitle", "DynRowSerial", "Dimension1"]

AUTUMN_SOWN_ROW_CODES = [
    "m5410",
    "m5411",
    "m5412",
    "m5413",
    "m5414",
    "m5415",
    "m5416",
    "m5417",
    "m5418",
    "m5422",
    "m5423",
    "m5424",
    "m5425",
    "m5429",
    "m5430",
]

AUTUMN_SOWN_LABELS = {
    "m5410": "Mezei leltár összesen",
    "m5411": "Őszi búza",
    "m5412": "Durumbúza",
    "m5413": "Rozs",
    "m5414": "Őszi árpa",
    "m5415": "Triticale",
    "m5416": "Repce",
    "m5417": "Őszi takarmánykeverék",
    "m5418": "Évelő pillangósok",
    "m5422": "Egyéb szántóföldi kultúra",
    "m5423": "Rét-legelő",
    "m5424": "Zöldségtermelés",
    "m5425": "Virág- és dísznövény termelés",
    "m5429": "Egyéb kertészeti termelés",
    "m5430": "Tavaszi vetések előkészítése",
}

STOCK_5C_TO_6B_PREFIX = {
    "m55": "m64",
    "m56": "m65",
    "m57": "m66",
    "m58": "m67",
}
STOCK_5C_PREFIXES = tuple(STOCK_5C_TO_6B_PREFIX)
STOCK_6B_PREFIXES = tuple(STOCK_5C_TO_6B_PREFIX.values())
STOCK_5C_CLOSING_COLUMN = "5"
STOCK_6B_CLOSING_COLUMN = "12"
ANIMAL_CLOSING_COLUMN = "12"


@dataclass(frozen=True)
class SheetTemplate:
    sheet_name: str
    title: str
    column_codes: list[str]
    column_labels: list[str]
    rows: list[dict[str, str]]


def _open_workbook(source: str | Path | bytes) -> Any:
    if isinstance(source, (str, Path)):
        return load_workbook(filename=source, read_only=True, data_only=True)
    return load_workbook(filename=BytesIO(source), read_only=True, data_only=True)


def _read_excel(source: str | Path | bytes) -> pd.DataFrame:
    archive = _open_zip_archive(source)
    shared_strings = _read_shared_strings(archive)
    sheet_xml = _read_first_sheet_xml(archive)
    rows = _parse_sheet_rows(sheet_xml, shared_strings)
    if not rows:
        return pd.DataFrame()

    headers = [_clean(value) for value in rows[0]]
    data_rows = rows[1:]

    frame = pd.DataFrame(data_rows, columns=headers)
    frame = frame.dropna(how="all").reset_index(drop=True)
    return frame


def _open_zip_archive(source: str | Path | bytes) -> zipfile.ZipFile:
    if isinstance(source, (str, Path)):
        return zipfile.ZipFile(source)
    return zipfile.ZipFile(BytesIO(source))


def _read_shared_strings(archive: zipfile.ZipFile) -> list[str]:
    if "xl/sharedStrings.xml" not in archive.namelist():
        return []

    root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
    values: list[str] = []
    for string_item in root.findall(".//{*}si"):
        chunks = [node.text or "" for node in string_item.findall(".//{*}t")]
        values.append("".join(chunks))
    return values


def _read_first_sheet_xml(archive: zipfile.ZipFile) -> bytes:
    workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
    sheets = workbook_root.findall(".//{*}sheet")
    if not sheets:
        raise ValueError("Nem talalhato munkalap az omlesztett fajlban.")

    relationship_id = None
    for key, value in sheets[0].attrib.items():
        if key.endswith("id"):
            relationship_id = value
            break

    rel_root = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    target = None
    for relation in rel_root.findall(".//{*}Relationship"):
        if relation.attrib.get("Id") == relationship_id:
            target = relation.attrib.get("Target", "")
            break

    if not target:
        raise ValueError("Nem talalhato a bulk munkalap XML-je.")

    normalized = target.lstrip("/")
    if not normalized.startswith("xl/"):
        normalized = f"xl/{normalized.lstrip('./')}"
    return archive.read(normalized)


def _parse_sheet_rows(sheet_xml: bytes, shared_strings: list[str]) -> list[list[str]]:
    root = ET.fromstring(sheet_xml)
    parsed_rows: list[list[str]] = []

    for row in root.findall(".//{*}sheetData/{*}row"):
        values_by_index: dict[int, str] = {}
        fallback_index = 0

        for cell in row.findall("{*}c"):
            column_index = _column_index_from_ref(cell.attrib.get("r")) or fallback_index
            values_by_index[column_index] = _parse_cell_value(cell, shared_strings)
            fallback_index = max(fallback_index, column_index + 1)

        if not values_by_index:
            continue

        width = max(values_by_index) + 1
        parsed_rows.append([values_by_index.get(index, "") for index in range(width)])

    return parsed_rows


def _column_index_from_ref(reference: str | None) -> int | None:
    if not reference:
        return None

    match = re.match(r"([A-Z]+)", reference.upper())
    if not match:
        return None

    letters = match.group(1)
    index = 0
    for letter in letters:
        index = index * 26 + (ord(letter) - ord("A") + 1)
    return index - 1


def _parse_cell_value(cell: ET.Element, shared_strings: list[str]) -> str:
    cell_type = cell.attrib.get("t", "")
    value_node = cell.find("{*}v")
    inline_node = cell.find("{*}is")

    if inline_node is not None:
        text_chunks = [node.text or "" for node in inline_node.findall(".//{*}t")]
        return "".join(text_chunks)

    raw_value = value_node.text if value_node is not None and value_node.text is not None else ""
    if cell_type == "s":
        try:
            return shared_strings[int(raw_value)]
        except (ValueError, IndexError):
            return raw_value
    return raw_value


def _clean(value: Any) -> str:
    if pd.isna(value):
        return ""
    if value is None:
        return ""
    text = str(value).strip()
    if text.casefold() == "nan":
        return ""
    return text


def _sort_key(value: str) -> tuple[int, str]:
    text = _clean(value)
    if text.isdigit():
        return (0, f"{int(text):08d}")
    return (1, text)


def load_template(source: str | Path | bytes) -> dict[str, Any]:
    workbook = _open_workbook(source)
    sheets: dict[str, SheetTemplate] = {}
    row_to_sheet: dict[str, str] = {}

    for worksheet in workbook.worksheets:
        rows = list(worksheet.iter_rows(values_only=True))
        if len(rows) < 4:
            continue

        title = _clean(rows[0][0]) or worksheet.title
        header_codes = [_clean(value) for value in rows[2]]
        header_labels = [_clean(value) for value in rows[3]]
        column_codes = header_codes[5:]
        column_labels = [
            label if label else f"Oszlop {code}"
            for code, label in zip(column_codes, header_labels[5:])
        ]

        template_rows: list[dict[str, str]] = []
        for raw_row in rows[4:]:
            row_code = _clean(raw_row[2]) if len(raw_row) > 2 else ""
            row_title = _clean(raw_row[3]) if len(raw_row) > 3 else ""
            dyn_row_serial = _clean(raw_row[4]) if len(raw_row) > 4 else ""
            if not row_code:
                continue
            template_rows.append(
                {
                    "row_code": row_code,
                    "row_title": row_title,
                    "dyn_row_serial": dyn_row_serial,
                }
            )
            row_to_sheet[row_code] = worksheet.title

        sheets[worksheet.title] = SheetTemplate(
            sheet_name=worksheet.title,
            title=title,
            column_codes=column_codes,
            column_labels=column_labels,
            rows=template_rows,
        )

    return {
        "sheets": sheets,
        "row_to_sheet": row_to_sheet,
    }


def load_bulk_data(source: str | Path | bytes) -> pd.DataFrame:
    dataframe = _read_excel(source)
    dataframe.columns = [str(column).strip().lower() for column in dataframe.columns]

    required = ["akod", "sor", "osz", "ertek", "dimenzio1", "dyn_row_serial"]
    missing = [column for column in required if column not in dataframe.columns]
    if missing:
        raise ValueError(f"Hiányzó oszlopok az ömlesztett fájlból: {', '.join(missing)}")

    prepared = dataframe[required].copy()
    for column in required:
        prepared[column] = prepared[column].map(_clean)

    prepared["sheet_name"] = prepared["sor"].map(lambda value: "")
    return prepared


def attach_sheet_names(dataframe: pd.DataFrame, row_to_sheet: dict[str, str]) -> pd.DataFrame:
    prepared = dataframe.copy()
    prepared["sheet_name"] = prepared["sor"].map(lambda value: row_to_sheet.get(value, ""))
    return prepared


def list_farms(dataframe: pd.DataFrame) -> list[str]:
    farms = sorted({value for value in dataframe["akod"].tolist() if value})
    return farms


def build_sheet_view(
    dataframe: pd.DataFrame,
    template: SheetTemplate,
    farm_code: str,
    filled_only: bool = False,
    closing_only: bool = False,
) -> pd.DataFrame:
    sheet_rows = dataframe[
        (dataframe["akod"] == farm_code) & (dataframe["sheet_name"] == template.sheet_name)
    ].copy()

    if sheet_rows.empty:
        return pd.DataFrame(columns=TECHNICAL_COLUMNS + template.column_labels)

    value_lookup = {
        (
            _clean(row.sor),
            _clean(row.dyn_row_serial),
            _clean(row.dimenzio1),
            _clean(row.osz),
        ): _clean(row.ertek)
        for row in sheet_rows.itertuples()
    }

    key_lookup: dict[str, set[tuple[str, str]]] = {}
    for row in sheet_rows.itertuples():
        key_lookup.setdefault(_clean(row.sor), set()).add(
            (_clean(row.dyn_row_serial), _clean(row.dimenzio1))
        )

    records: list[dict[str, str]] = []
    for template_row in template.rows:
        row_code = template_row["row_code"]
        row_title = template_row["row_title"]
        default_dyn = template_row["dyn_row_serial"] or "0"
        keys = sorted(key_lookup.get(row_code, {(default_dyn, "")}), key=lambda item: (_sort_key(item[0]), item[1]))

        for dyn_row_serial, dimension1 in keys:
            record = {
                "FarmCode": farm_code,
                "RowCode": row_code,
                "RowTitle": row_title,
                "DynRowSerial": dyn_row_serial or default_dyn,
                "Dimension1": dimension1,
            }
            has_value = False
            for code, label in zip(template.column_codes, template.column_labels):
                value = value_lookup.get((row_code, dyn_row_serial, dimension1, code), "")
                record[label] = value
                has_value = has_value or _has_meaningful_value(value)

            if filled_only and not has_value:
                continue
            if closing_only and not _looks_like_closing_row(row_title):
                continue
            records.append(record)

    return pd.DataFrame(records)


def build_land_view(
    dataframe: pd.DataFrame,
    template: SheetTemplate,
    farm_code: str,
) -> pd.DataFrame:
    """Build the targeted t1_a export from rows with meaningful values.

    The output follows the visible general-view column order, but excludes
    hidden technical columns and the trailing empty template header. Every
    template row with at least one non-empty, non-zero value is retained,
    including calculated and summary rows.
    """

    value_columns = [
        (code, label)
        for code, label in zip(template.column_codes, template.column_labels)
        if _clean(code) and _clean(label)
    ]
    columns = ["FarmCode", "RowCode", "RowTitle"] + [
        label for _, label in value_columns
    ]

    visible_view = build_sheet_view(
        dataframe=dataframe,
        template=template,
        farm_code=farm_code,
        filled_only=True,
        closing_only=False,
    )
    if visible_view.empty:
        return pd.DataFrame(columns=columns)

    export_view = pd.DataFrame(
        {
            "FarmCode": visible_view["FarmCode"],
            "RowCode": visible_view["RowCode"],
            "RowTitle": visible_view["RowTitle"],
        }
    )
    for _, label in value_columns:
        export_view[label] = visible_view[label].map(_parse_export_number)

    return export_view[columns].reset_index(drop=True)


def _parse_export_number(value: Any) -> float | None:
    text = _clean(value)
    if not text:
        return None

    if text.startswith("="):
        text = text[1:].strip()

    compact = text.replace(" ", "")
    if "," in compact and "." in compact:
        if compact.rfind(",") > compact.rfind("."):
            compact = compact.replace(".", "").replace(",", ".")
        else:
            compact = compact.replace(",", "")
    else:
        compact = compact.replace(",", ".")

    try:
        number = float(compact)
    except ValueError:
        return None

    if number == 0:
        return None
    return round(number, 2)


def build_autumn_sown_view(
    dataframe: pd.DataFrame,
    template: SheetTemplate,
    farm_code: str,
) -> pd.DataFrame:
    """Build the first targeted export: non-zero autumn-sown area rows from 5C."""

    columns = ["Üzemkód", "RowCode", "Sor megnevezése", "Érték"]
    sheet_rows = dataframe[
        (dataframe["akod"] == farm_code)
        & (dataframe["sheet_name"] == template.sheet_name)
        & (dataframe["osz"] == "5")
    ]

    template_rows = {row["row_code"]: row for row in template.rows}
    records: list[dict[str, Any]] = []
    wheat_values: dict[str, float] = {}

    for row_code in AUTUMN_SOWN_ROW_CODES:
        template_row = template_rows.get(row_code)
        if template_row is None:
            continue

        source_rows = sheet_rows[sheet_rows["sor"] == row_code]
        values = [
            number
            for number in (_parse_export_number(value) for value in source_rows["ertek"].tolist())
            if number is not None
        ]
        if values:
            value = round(sum(values), 2)
            record = {
                "Üzemkód": farm_code,
                "RowCode": row_code,
                "Sor megnevezése": AUTUMN_SOWN_LABELS.get(
                    row_code,
                    template_row["row_title"],
                ),
                "Érték": value,
            }
            records.append(record)

            if row_code in {"m5411", "m5412"}:
                wheat_values[row_code] = value

        if row_code == "m5412" and wheat_values:
            records.append(
                {
                    "Üzemkód": farm_code,
                    "RowCode": "",
                    "Sor megnevezése": "Búza",
                    "Érték": round(sum(wheat_values.values()), 2),
                }
            )

    return pd.DataFrame(records, columns=columns)


def _is_stock_summary_row(row_title: str) -> bool:
    normalized = _clean(row_title).casefold()
    return "összesen" in normalized or "mindösszesen" in normalized


def _stock_code_is_in_family(row_code: str, prefixes: tuple[str, ...]) -> bool:
    return any(_clean(row_code).startswith(prefix) for prefix in prefixes)


def _map_stock_5c_code_to_6b(row_code: str) -> str | None:
    cleaned = _clean(row_code)
    for source_prefix, target_prefix in STOCK_5C_TO_6B_PREFIX.items():
        if cleaned.startswith(source_prefix):
            return f"{target_prefix}{cleaned[len(source_prefix):]}"
    return None


def _stock_closing_lookup(
    dataframe: pd.DataFrame,
    farm_code: str,
) -> dict[tuple[str, str], float]:
    source_rows = dataframe[
        (dataframe["akod"] == farm_code)
        & dataframe["osz"].isin(
            {STOCK_5C_CLOSING_COLUMN, STOCK_6B_CLOSING_COLUMN}
        )
    ].copy()
    if source_rows.empty:
        return {}

    source_rows["_parsed_value"] = (
        source_rows["ertek"].map(_parse_export_number).fillna(0.0)
    )
    grouped = source_rows.groupby(["sor", "osz"])["_parsed_value"].sum()
    return {
        (_clean(row_code), _clean(column_code)): round(float(value), 2)
        for (row_code, column_code), value in grouped.items()
    }


def build_stock_view(
    dataframe: pd.DataFrame,
    template_5c: SheetTemplate,
    template_6b: SheetTemplate,
    farm_code: str,
) -> pd.DataFrame:
    """Build the targeted stock export from 5C and 6B closing quantities.

    The 5C-to-6B relationship is derived from the intentional code-family
    rule. Summary rows are excluded before matching. A detailed 5C row without
    its derived 6B counterpart is treated as a template error rather than
    being silently omitted.
    """

    columns = [
        "Partner azonosító",
        "5C RowCode",
        "5C megnevezés",
        "5C záróérték",
        "6B RowCode",
        "6B megnevezés",
        "6B zárókészlet",
        "Saját készlet",
        "Vásárolt készlet",
        "Ellenőrzés",
    ]

    rows_5c = {
        row["row_code"]: row
        for row in template_5c.rows
        if _stock_code_is_in_family(row["row_code"], STOCK_5C_PREFIXES)
        and not _is_stock_summary_row(row["row_title"])
    }
    rows_6b = {
        row["row_code"]: row
        for row in template_6b.rows
        if _stock_code_is_in_family(row["row_code"], STOCK_6B_PREFIXES)
        and not _is_stock_summary_row(row["row_title"])
    }

    missing_6b_pairs = [
        row_code
        for row_code in rows_5c
        if _map_stock_5c_code_to_6b(row_code) not in rows_6b
    ]
    if missing_6b_pairs:
        missing_text = ", ".join(missing_6b_pairs)
        raise ValueError(
            "A készlet-exporthoz részletes 5C sorhoz nem található 6B pár "
            f"a kódszabály alapján: {missing_text}."
        )

    closing_values = _stock_closing_lookup(dataframe, farm_code)
    records: list[dict[str, Any]] = []
    paired_6b_codes: set[str] = set()

    for row_code_5c, row_5c in rows_5c.items():
        row_code_6b = _map_stock_5c_code_to_6b(row_code_5c)
        if row_code_6b is None:
            continue

        row_6b = rows_6b[row_code_6b]
        paired_6b_codes.add(row_code_6b)
        own_source = closing_values.get(
            (row_code_5c, STOCK_5C_CLOSING_COLUMN),
            0.0,
        )
        total_stock = closing_values.get(
            (row_code_6b, STOCK_6B_CLOSING_COLUMN),
            0.0,
        )

        if own_source == 0 and total_stock == 0:
            continue

        if total_stock < own_source:
            own_stock = 0.0
            purchased_stock = 0.0
            status = "HIBA: a 6B zárókészlet kisebb az 5C saját készletnél"
        else:
            own_stock = own_source
            purchased_stock = round(total_stock - own_source, 2)
            status = "Rendben"

        records.append(
            {
                "Partner azonosító": farm_code,
                "5C RowCode": row_code_5c,
                "5C megnevezés": row_5c["row_title"],
                "5C záróérték": own_source,
                "6B RowCode": row_code_6b,
                "6B megnevezés": row_6b["row_title"],
                "6B zárókészlet": total_stock,
                "Saját készlet": own_stock,
                "Vásárolt készlet": purchased_stock,
                "Ellenőrzés": status,
            }
        )

    for row_code_6b, row_6b in rows_6b.items():
        if row_code_6b in paired_6b_codes:
            continue

        total_stock = closing_values.get(
            (row_code_6b, STOCK_6B_CLOSING_COLUMN),
            0.0,
        )
        if total_stock == 0:
            continue

        records.append(
            {
                "Partner azonosító": farm_code,
                "5C RowCode": None,
                "5C megnevezés": None,
                "5C záróérték": None,
                "6B RowCode": row_code_6b,
                "6B megnevezés": row_6b["row_title"],
                "6B zárókészlet": total_stock,
                "Saját készlet": 0.0,
                "Vásárolt készlet": total_stock,
                "Ellenőrzés": "Rendben – csak 6B-ben szerepel",
            }
        )

    return pd.DataFrame(records, columns=columns)


def _animal_unit(row_code: str, row_title: str) -> str:
    normalized_title = _clean(row_title).casefold()
    if "t-ban" in normalized_title:
        return "t"
    if _clean(row_code) == "m6323":
        return "család"
    return "db"


def build_animal_view(
    dataframe: pd.DataFrame,
    template: SheetTemplate,
    farm_code: str,
) -> pd.DataFrame:
    """Build the targeted 6A animal export from non-zero closing values.

    The 6A sheet contains movements in several osz columns. Only osz=12
    (closing stock) is relevant here. Summary rows are intentionally omitted,
    while non-zero companion rows labelled ``előző sor t-ban`` are retained
    because they contain the closing weight of the preceding animal row.
    """

    columns = [
        "Üzemkód",
        "RowCode",
        "Sor megnevezése",
        "Mértékegység",
        "Záróérték",
    ]
    sheet_rows = dataframe[
        (dataframe["akod"] == farm_code)
        & (dataframe["sheet_name"] == template.sheet_name)
        & (dataframe["osz"] == ANIMAL_CLOSING_COLUMN)
    ].copy()
    if sheet_rows.empty:
        return pd.DataFrame(columns=columns)

    closing_values = (
        sheet_rows.assign(_parsed_value=sheet_rows["ertek"].map(_parse_export_number))
        .dropna(subset=["_parsed_value"])
        .groupby("sor")["_parsed_value"]
        .sum()
    )

    records: list[dict[str, Any]] = []
    for template_row in template.rows:
        row_code = template_row["row_code"]
        row_title = template_row["row_title"]
        if _is_stock_summary_row(row_title):
            continue

        value = closing_values.get(row_code)
        if value is None or abs(float(value)) < 1e-12:
            continue

        records.append(
            {
                "Üzemkód": farm_code,
                "RowCode": row_code,
                "Sor megnevezése": row_title,
                "Mértékegység": _animal_unit(row_code, row_title),
                "Záróérték": round(float(value), 2),
            }
        )

    return pd.DataFrame(records, columns=columns)


def _looks_like_closing_row(text: str) -> bool:
    normalized = _clean(text).casefold()
    keywords = [
        "záró",
        "záróérték",
        "záró érték",
        "záró-",
        "összesen",
        "mindösszesen",
        "év végi",
        "állomány",
    ]
    return any(keyword in normalized for keyword in keywords)


def _has_meaningful_value(value: str) -> bool:
    text = _clean(value)
    if not text:
        return False

    if text.startswith("="):
        text = text[1:].strip()

    return re.fullmatch(r"0+(?:[.,]0+)?", text) is None


def normalize_display_value(value: Any) -> str:
    text = _clean(value)
    if not text:
        return ""

    plain = text
    if plain.startswith("="):
        plain = plain[1:].strip()

    if re.fullmatch(r"0+(?:[.,]0+)?", plain):
        return ""

    return text


def summarize_template(template_bundle: dict[str, Any]) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for sheet in template_bundle["sheets"].values():
        rows.append(
            {
                "Munkalap": sheet.sheet_name,
                "Cím": sheet.title,
                "Sorok száma": len(sheet.rows),
                "Értékoszlopok": len(sheet.column_codes),
            }
        )
    return pd.DataFrame(rows).sort_values("Munkalap").reset_index(drop=True)
